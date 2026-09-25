from docx import Document
from docx.shared import Inches
from pathlib import Path

report_dir = Path(r'd:\student-portfolio\reports')
image_path = report_dir / 'practical6-screenshot.png'
doc_path = report_dir / 'Practical_6_Full_Stack_Integration.docx'

if not image_path.exists():
    raise FileNotFoundError(f'Missing screenshot: {image_path}')

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

p = doc.add_paragraph()
run = p.add_run('Practical 6: Full-Stack Integration with React, Node/Express and MongoDB')
run.bold = True
run.font.size = 24

section_headers = [
    ('1. Objective', 'To wire the React frontend to the Node/Express/MongoDB backend into a fully functional full-stack application.'),
    ('2. Prerequisites', '• Practicals 1–3 completed with a working React UI, routing, and API consumption patterns.\n• Practicals 4–5 completed with a working Node/Express/MongoDB backend with CRUD operations.'),
    ('3. Theory Concepts', 'Architecture and full-stack request flow:\nReact Frontend (localhost:5173)\n        |\n     fetch / axios\n        |\nExpress Backend (localhost:5000)\n        |\n      Mongoose\n        |\nMongoDB Database\nFlow: UI Form → POST /tasks → MongoDB → UI re-fetches → updated list rendered.'),
    ('4. Problem Definition', '• Connect the React Task UI to the Node + MongoDB backend built in Practicals 4 and 5.\n• Replace the Practical 3 GitHub repository data with task data from the student\'s own backend.\n• Support creating, viewing, updating, and deleting tasks from the React interface through API calls.\n• Persist all data in MongoDB and confirm persistence by refreshing the browser.\n• Handle loading and error states for every API interaction, including POST, PUT, and DELETE.'),
    ('5. Key Questions / Analysis', '• What changes are required on the backend (CORS) to allow the React development server to call the Express API?\n• Why should the UI re-fetch or update local state after a successful POST/PUT/DELETE?\n• What is the risk of not handling errors on write operations in the same way as read operations?'),
    ('6. Lab Session Step-by-Step', '1. Open the backend project and install CORS using: npm install cors.\n2. Import CORS in the Express server and enable it before the route definitions.\n3. Create a central React api.js file containing the backend base URL.\n4. Create getTasks() to call GET /tasks.\n5. Replace the Practical 3 GitHub fetch logic with the application\'s own /tasks endpoint.\n6. Create a React task form that sends POST /tasks and updates local state after success.\n7. Implement update using PUT /tasks/:id and synchronize the task list after success.\n8. Implement delete using DELETE /tasks/:id and synchronize the task list after success.\n9. Add loading and error handling around every API operation.\n10. Run both development servers simultaneously: React on port 5173 and Express on port 5000.\n11. Test the complete flow: create → view → update → delete.\n12. Refresh the browser and confirm that persisted task data is loaded again from MongoDB.'),
    ('7. Backend CORS Configuration', 'npm install cors\n\nconst cors = require(\'cors\');\n\napp.use(cors());\n\nCORS should be configured before the route definitions so the React development server can communicate with the Express API.'),
    ('8. React API Configuration', '// src/api.js\n\nconst BASE_URL = \'http://localhost:5000\';\n\nexport const getTasks = () =>\n\n  fetch(`${BASE_URL}/tasks`)\n\n    .then(res => res.json());'),
]

for title, content in section_headers:
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = 16
    doc.add_paragraph(content)

p = doc.add_paragraph()
run = p.add_run('Screenshot of the React Task Manager UI')
run.bold = True
run.font.size = 14

doc.add_picture(str(image_path), width=Inches(7.5))

doc.add_paragraph()
conclusion = doc.add_paragraph('Conclusion')
conclusion.runs[0].bold = True
conclusion.runs[0].font.size = 16

doc.add_paragraph('This practical demonstrates full-stack integration by connecting the React frontend to an Express API backed by MongoDB. The application supports creating, reading, updating, and deleting tasks while handling loading states, validation, and errors in a real-world client-server flow.')

doc.save(doc_path)
print(f'Created: {doc_path}')
print(f'Exists: {doc_path.exists()}')
print(f'Bytes: {doc_path.stat().st_size if doc_path.exists() else 0}')
